from pathlib import Path
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def export_to_excel(data: Dict[str, Any], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    headers = data.get("headers", [])
    rows = data.get("rows", [])
    
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(value, str) and value.startswith("="):
                cell.value = value
            else:
                cell.value = value
    
    for col_idx, header in enumerate(headers, start=1):
        max_length = len(str(header))
        for row_idx in range(2, len(rows) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)